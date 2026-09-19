from pptx import Presentation

prs = Presentation()

layout = prs.slide_layouts[0]

slide = prs.slides.add_slide(layout)

prs.save("presentation.pptx")